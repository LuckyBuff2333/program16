# -*- coding: utf-8 -*-
import openpyxl
wb = openpyxl.load_workbook(r'e:\program16\触发时间修复.xlsx')
ws = wb.active
print(f"Rows: {ws.max_row}, Cols: {ws.max_column}")
print(f"Headers: {[c.value for c in ws[1]]}")
print("Sample rows:")
for row in ws.iter_rows(min_row=2, max_row=6):
    print(f"  {[c.value for c in row]}")
