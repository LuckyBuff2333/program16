"""触发时间管理路由：批量读取、更新、重新提取、云端缓存批量操作"""
import asyncio
import os

from fastapi import APIRouter, Request, UploadFile, File
from fastapi.responses import StreamingResponse

from src.config import PROJECT_ROOT, load_config, setup_logger
from src.core.attachment_guard import begin_batch as _begin_attach_batch
from src.api.routes.shared import run_blocking as _run_blocking
from src.services.trigger_ops import remember_ref_time as _remember_ref_time
from src.services.server_legacy import _diag_gmlogger_times as _gm_times_of
from src.schemas.common import ApiResponse
from src.schemas.trigger import (
    BugidListRequest, UpdateTriggerTimeRequest, BatchUpdateRequest,
    ReExtractScanRequest, CacheBatchScanRequest, CacheBatchExtractRequest,
)

logger = setup_logger("web.trigger")
router = APIRouter()


@router.post("/api/trigger_time/batch_read", response_model=ApiResponse)
async def batch_read_trigger_times(req: BugidListRequest):
    """批量读取指定 jira 号的当前触发时间"""
    from src.services.trigger_ops import (
        _load_trigger_times_from_cloud, _save_trigger_time_to_csv,
        _batch_save_trigger_times_to_cloud,
    )
    from src.api.routes.shared import ok as _ok, fail as _fail
    bugids = req.bugids
    if not bugids:
        return _fail("请输入 jira 号列表")
    local_times = _load_trigger_times_from_cloud()
    result, extracted_count, batch_to_cloud = [], 0, {}
    batch_refs = {}
    for bugid in bugids:
        bugid = bugid.strip()
        if not bugid:
            continue
        current_time = local_times.get(bugid, "")
        if not current_time:
            try:
                from src.clients import jira_client as _jc
                loop = asyncio.get_event_loop()
                issue = await loop.run_in_executor(None, _jc.fetch_issue, bugid)
                tt = await loop.run_in_executor(None, _jc.extract_trigger_time_from_issue, issue)
                if tt:
                    current_time = tt
                    _save_trigger_time_to_csv(bugid, tt)
                    batch_to_cloud[bugid] = tt
                    try:
                        from src.services.server_legacy import _diag_gmlogger_times as _gm_times
                        _gm = _gm_times((issue.get("fields") or {}).get("attachment") or []) or []
                        if _gm:
                            _d = _gm[0][1]
                            batch_refs[bugid] = (_d.strftime("%Y-%m-%d %H:%M:%S")
                                                 if hasattr(_d, "strftime") else str(_d)[:19])
                    except Exception:
                        pass
                    extracted_count += 1
            except Exception as e:
                logger.warning("批量读取触发时间-提取失败 %s: %s", bugid, e)
        result.append({"bugid": bugid, "current_time": current_time})
    if batch_to_cloud:
        _batch_save_trigger_times_to_cloud(
            batch_to_cloud, {k: "时间对比(读取)" for k in batch_to_cloud},
            {k: v for k, v in batch_refs.items() if v})
    hit_count = sum(1 for r in result if r["current_time"])
    return _ok({"rows": result, "total": len(result), "hit": hit_count},
               f"查询完成：{len(result)} 个 Jira 号，{hit_count} 个有记录（本地命中 {hit_count - extracted_count}，新提取 {extracted_count}）")


@router.post("/api/trigger_time/update", response_model=ApiResponse)
def update_trigger_time(req: UpdateTriggerTimeRequest):
    """修改指定bugid的触发时间"""
    from src.api.routes.shared import (
        ok as _ok,
        fail as _fail,
    )
    from src.services.report_ops import (
        _load_diagnose_cache,
        _save_diagnose_cache,
    )
    from src.services.trigger_ops import _save_trigger_time_to_cloud
    bugid = req.bugid.strip()
    new_time = req.new_time.strip()
    if not bugid:
        return _fail("请输入 bugid/jira号")
    if not new_time:
        return _fail("请输入新的触发时间")
    try:
        _save_trigger_time_to_cloud(bugid, new_time, source="人工修正")
        cached = _load_diagnose_cache(bugid)
        if cached:
            cached["verify_status"] = "corrected"
            cached["suggested_time"] = new_time
            _save_diagnose_cache(bugid, cached)
        return _ok({"bugid": bugid, "new_time": new_time, "updated": True},
                   f"已更新 {bugid} 的触发时间为 {new_time}（云端+本地+诊断缓存）")
    except Exception as e:
        logger.warning("更新触发时间失败 %s: %s", bugid, e)
        return _fail(f"更新失败: {e}")


@router.post("/api/trigger_time/batch_update", response_model=ApiResponse)
def batch_update_trigger_time(req: BatchUpdateRequest):
    """批量应用修正时间"""
    from src.services.trigger_ops import _batch_save_trigger_times_to_cloud
    from src.services.report_ops import _load_diagnose_cache, _save_diagnose_cache
    from src.api.routes.shared import ok as _ok, fail as _fail
    items = req.items
    if not items:
        return _fail("待应用列表为空")
    valid_items = {}
    for item in items:
        bugid = item.bugid.strip()
        new_time = item.new_time.strip()
        if not bugid or not new_time:
            continue
        valid_items[bugid] = new_time
    if not valid_items:
        return _fail("待应用列表为空")
    cloud_ok = _batch_save_trigger_times_to_cloud(
        valid_items, {k: "人工修正(批量)" for k in valid_items})
    if not cloud_ok:
        return _fail("云端写入失败")
    cache_updated = 0
    for bugid, new_time in valid_items.items():
        cached = _load_diagnose_cache(bugid)
        if cached:
            cached["verify_status"] = "corrected"
            cached["suggested_time"] = new_time
            _save_diagnose_cache(bugid, cached)
            cache_updated += 1
    success = len(valid_items)
    msg = f"批量应用完成：成功 {success} 条"
    return _ok({"success": success, "failed": 0}, msg)


@router.post("/api/trigger_time/re_extract_scan", response_model=ApiResponse)
async def re_extract_scan(req: ReExtractScanRequest):
    """扫描云端触发时间记录，找出触发时间为空的候选项"""
    from src.api.routes.shared import (
        bitable_text as _bitable_text,
        ok as _ok,
        fail as _fail,
    )
    from src.services.trigger_ops import _load_all_cloud_trigger_records
    cache_only = req.cache_only
    try:
        loop = asyncio.get_event_loop()
        all_records = await loop.run_in_executor(None, _load_all_cloud_trigger_records)
        empty_records = [r for r in all_records if not r["trigger_time"]]
        if not empty_records:
            return _ok({"candidates": [], "total_empty": 0, "total_cloud": len(all_records)},
                       "云端所有记录均已提取到触发时间")
        if cache_only:
            candidates = [{"bugid": r["jira_no"], "jira_status": "", "table_id": r["table_id"], "record_id": r["record_id"]} for r in empty_records]
            return _ok({"candidates": candidates, "total_empty": len(empty_records), "total_cloud": len(all_records)},
                       f"缓存扫描完成: 云端共 {len(all_records)} 条，触发时间为空 {len(candidates)} 条")
        from src.clients import jira_client as _jc
        candidates = []
        for i, rec in enumerate(empty_records):
            if await request.is_disconnected():
                break
            await asyncio.sleep(0)
            try:
                issue = await loop.run_in_executor(None, _jc.fetch_issue, rec["jira_no"])
                status = ((issue.get("fields") or {}).get("status") or {}).get("name", "")
                if status.lower() != "closed":
                    candidates.append({"bugid": rec["jira_no"], "jira_status": status, "table_id": rec["table_id"], "record_id": rec["record_id"]})
            except Exception as e:
                candidates.append({"bugid": rec["jira_no"], "jira_status": f"查询失败: {e}", "table_id": rec["table_id"], "record_id": rec["record_id"]})
        return _ok({"candidates": candidates, "total_empty": len(empty_records), "total_cloud": len(all_records)},
                    f"扫描完成: 云端共 {len(all_records)} 条，空触发时间 {len(empty_records)} 条，可重新提取（非 Closed）{len(candidates)} 条")
    except Exception as e:
        logger.error("重新提取扫描失败: %s", e)
        return _fail(f"扫描失败: {e}")


@router.post("/api/trigger_time/re_extract", response_model=ApiResponse)
async def re_extract(req: BugidListRequest):
    """对指定的 bugid 列表重新提取触发时间（视频优先）"""
    from src.services.server_legacy import _diag_extract_time_from_video_simple, _diag_extract_time_from_images
    from src.services.trigger_ops import _save_trigger_time_to_cloud
    from src.api.routes.shared import ok as _ok, fail as _fail
    from src.clients import jira_client as _jc
    from src.config import load_config
    bugids = req.bugids
    if not bugids:
        return _fail("待重新提取列表为空")
    try:
        loop = asyncio.get_event_loop()
        results = []
        for i, bugid in enumerate(bugids):
            if await request.is_disconnected():
                break
            await asyncio.sleep(0)
            try:
                issue = await loop.run_in_executor(None, _jc.fetch_issue, bugid)
                tt = await loop.run_in_executor(None, _diag_extract_time_from_video_simple, issue)
                if not tt:
                    tt = await loop.run_in_executor(None, lambda iss=issue: _jc.extract_trigger_time_from_issue(iss, skip_attachment=True))
                # 图片 OCR 兑底
                if not tt:
                    cfg_api = load_config().get("jira_api", {})
                    img_h, img_tok = {}, cfg_api.get("token", "")
                    if img_tok:
                        img_h["Authorization"] = f"Bearer {img_tok}"
                    img_to = max(int(cfg_api.get("timeout", 30)), 120)
                    img_result = await loop.run_in_executor(None, _diag_extract_time_from_images, issue, img_h, img_to, bugid)
                    img_tt = img_result[0] if isinstance(img_result, tuple) else img_result
                    if img_tt:
                        tt = img_tt
                # 附件文件名最终兑底
                if not tt:
                    tt = await loop.run_in_executor(None, _jc.extract_attachment_time, issue)
                status = ((issue.get("fields") or {}).get("status") or {}).get("name", "")
                if tt:
                    await loop.run_in_executor(None, _save_trigger_time_to_cloud, bugid, tt)
                    results.append({"bugid": bugid, "trigger_time": tt, "status": status, "success": True})
                else:
                    results.append({"bugid": bugid, "trigger_time": "", "status": status, "success": False, "error": "未提取到时间"})
            except Exception as e:
                results.append({"bugid": bugid, "trigger_time": "", "success": False, "error": str(e)})
        success_count = sum(1 for r in results if r["success"])
        return _ok({"results": results, "success_count": success_count, "total": len(bugids)},
                    f"重新提取完成: {success_count}/{len(bugids)} 成功")
    except Exception as e:
        logger.error("重新提取失败: %s", e)
        return _fail(f"重新提取失败: {e}")


@router.post("/api/trigger_time/cache_batch_scan", response_model=ApiResponse)
async def cache_batch_scan(req: CacheBatchScanRequest):
    """云端缓存批量更新：扫描需要重新提取触发时间的候选记录"""
    from src.api.routes.shared import (
        ok as _ok,
        fail as _fail,
    )
    from src.services.trigger_ops import _load_all_cloud_trigger_records
    from src.services import trigger_service as svc
    exclude_bitable_success = req.exclude_bitable_success
    exclude_corrected = req.exclude_corrected
    try:
        loop = asyncio.get_event_loop()
        all_records = await loop.run_in_executor(None, _load_all_cloud_trigger_records)
        result = svc.scan_candidates(all_records, exclude_bitable_success, exclude_corrected)
        return _ok({
            "candidates": result["candidates"],
            "total_cloud": result["total_cloud"],
            "total_cached": result["total_cached"],
            "bitable_success_count": result["bitable_success_count"],
            "corrected_count": result["corrected_count"],
        }, result["message"])
    except Exception as e:
        logger.error("云端缓存批量扫描失败: %s", e)
        return _fail(f"扫描失败: {e}")


@router.post("/api/trigger_time/cache_batch_extract/stop", response_model=ApiResponse)
def stop_cache_batch_extract():
    """触发云端缓存批量提取停止信号"""
    from src.api.routes.shared import cache_batch_stop_event as _cache_batch_stop_event, ok
    _cache_batch_stop_event.set()
    logger.info("收到云端缓存批量提取停止信号")
    return ok(None, "停止信号已发送，当前条目完成后停止")


@router.post("/api/trigger_time/cache_batch_load_progress", response_model=ApiResponse)
def cache_batch_load_progress():
    """加载本地缓存批量提取进度"""
    from src.api.routes.shared import ok as _ok
    from src.services.trigger_ops import _load_cache_batch_progress
    data = _load_cache_batch_progress()
    return _ok({"progress": data}, f"本地进度: {len(data)} 条记录")


@router.post("/api/trigger_time/cache_batch_clear_progress", response_model=ApiResponse)
def cache_batch_clear_progress():
    """清除本地缓存批量提取进度"""
    from src.api.routes.shared import (
        ok as _ok,
        fail as _fail,
    )
    from src.services.trigger_ops import _CACHE_BATCH_PROGRESS_FILE
    try:
        if os.path.exists(_CACHE_BATCH_PROGRESS_FILE):
            os.remove(_CACHE_BATCH_PROGRESS_FILE)
        return _ok({}, "本地进度已清除")
    except Exception as e:
        return _fail(f"清除失败: {e}")


@router.post("/api/trigger_time/cache_batch_extract")
def cache_batch_extract(req: CacheBatchExtractRequest, request: Request):
    """云端缓存批量更新：SSE 流式推送进度，重新提取触发时间"""
    from src.api.routes.shared import cache_batch_stop_event as _cache_batch_stop_event, fail as _fail
    from src.services.trigger_ops import (
        _load_cache_batch_progress, _save_cache_batch_progress,
        _save_trigger_time_to_cloud,
    )
    from src.services.server_legacy import _diag_extract_time_from_video_simple, _VIDEO_SUFFIXES
    from src.services import trigger_service as svc
    bugids = req.bugids
    old_times = req.old_times
    video_first = req.video_first
    if not bugids:
        return _fail("待提取列表为空")

    gen = svc.generate_extract_events(
        bugids=bugids,
        old_times=old_times,
        video_first=video_first,
        stop_event=_cache_batch_stop_event,
        progress_loader=_load_cache_batch_progress,
        progress_saver=_save_cache_batch_progress,
        video_extractor=_diag_extract_time_from_video_simple,
        cloud_saver=_save_trigger_time_to_cloud,
        video_suffixes=_VIDEO_SUFFIXES,
        request=request,
    )
    # 注册作业：/api/jobs 可查进度，/api/jobs/{id}/stop 可单独停止本作业
    from src.core.job_registry import JOBS as _JOBS
    from src.api.routes.jobs import job_stream as _job_stream
    from src.api.routes.shared import cache_batch_stop_event as _job_stop
    _job = _JOBS.create("cache_batch", "缓存批处理作业")
    return StreamingResponse(_job_stream(_job, gen, _job_stop), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@router.post("/api/trigger_time/rebalance", response_model=ApiResponse)
def rebalance_trigger_time_tables_endpoint():
    """再平衡触发时间表：把超过500的表中多余记录搬运到不足500的表中"""
    from src.services.trigger_ops import rebalance_trigger_time_tables
    from src.api.routes.shared import ok as _ok, fail as _fail
    try:
        result = rebalance_trigger_time_tables()
        if result["moved"] > 0:
            return _ok(result, f"再平衡完成：搬运 {result['moved']} 条记录\n" + "\n".join(result["details"]))
        else:
            return _ok(result, result["details"][0] if result["details"] else "无需再平衡")
    except Exception as e:
        logger.error("触发时间表再平衡失败: %s", e, exc_info=True)
        return _fail(f"再平衡失败: {e}")


@router.post("/api/trigger/batch_extract_time")
async def batch_extract_time(request: Request):
    """批量提取触发时间（SSE 流式）：支持 JQL 查询或直接传入 bugid 列表

    提取链路：视频OCR → 评论/描述 → 图片OCR
    并行数 2，缓冲 20 条批量写入云端
    """
    import json as _json
    import concurrent.futures
    from src.api.routes.shared import batch_extract_stop_event as _stop_event, fail as _fail
    from src.clients import jira_client as _jc
    from src.services.server_legacy import (
        _diag_extract_time_from_video_simple, _diag_extract_time_from_images,
        _VIDEO_SUFFIXES,
    )
    from src.services.trigger_ops import (
        _load_trigger_times_from_cloud,
        _batch_save_trigger_times_to_cloud,
    )
    body = await request.json() or {}
    jql = str(body.get("jql", "")).strip()
    bugids = body.get("bugids") or []
    skip_cached = body.get("skip_cached", False)
    filter_ai_preliminary = body.get("filter_ai_preliminary", False)
    # ── 解析输入 ──
    if jql:
        try:
            issues = await _run_blocking(_jc.search_issues, jql, max_results=500)
            bugids = [(i.get("key") or "").strip() for i in issues if i.get("key")]
            bugids = [b for b in bugids if b]
        except Exception as e:
            return _fail(f"JQL 查询失败: {e}")
    if not bugids:
        return _fail("无待提取的 Jira 号（JQL 未匹配或 bugids 为空）")

    _stop_event.clear()
    _begin_attach_batch()  # 新批次：重置附件下载预算

    def event_generator():
        # 跳过已有缓存时，预加载云端缓存
        cloud_cache = {}
        if skip_cached:
            cloud_cache = _load_trigger_times_from_cloud()
        # 分离需要提取的和可跳过的（缓存跳过）
        to_extract = []
        cache_skip_count = 0
        for b in bugids:
            if skip_cached and cloud_cache.get(b):
                cache_skip_count += 1
                continue
            to_extract.append(b)

        # 过滤AI初步分析结果：预先检查 Jira 自定义字段
        ai_skip_count = 0
        if filter_ai_preliminary and to_extract:
            import concurrent.futures as _cf
            _AI_INIT_FIELD_ID = "customfield_13714"
            ai_has = set()

            def _check_ai(bid):
                try:
                    iss = _jc.fetch_issue(bid)
                    f = (iss.get("fields") or {}).get(_AI_INIT_FIELD_ID)
                    if f and str(f).strip() and str(f).strip().lower() not in ("none", "null"):
                        return bid
                except Exception:
                    pass
                return None

            with _cf.ThreadPoolExecutor(max_workers=10) as pool:
                for result in pool.map(_check_ai, to_extract):
                    if result:
                        ai_has.add(result)
            ai_skip_count = len(ai_has)
            if ai_has:
                to_extract = [b for b in to_extract if b not in ai_has]

        total = len(to_extract)
        cfg = load_config().get("jira_api", {})
        img_tok = cfg.get("token", "")
        img_to = max(int(cfg.get("timeout", 30)), 120)
        _img_headers = {}
        if img_tok:
            _img_headers["Authorization"] = f"Bearer {img_tok}"

        start_evt = {'type': 'start', 'total': total}
        skip_parts = []
        if ai_skip_count:
            start_evt['ai_skipped'] = ai_skip_count
            skip_parts.append(f'{ai_skip_count} 个已有AI分析')
        if cache_skip_count:
            start_evt['cache_skipped'] = cache_skip_count
            skip_parts.append(f'{cache_skip_count} 个已有缓存')
        if skip_parts:
            start_evt['message'] = f'跳过 {"、".join(skip_parts)}，实际提取 {total} 个'
        yield f"data: {_json.dumps(start_evt, ensure_ascii=False)}\n\n"

        done_count = 0
        hit_count = 0
        batch_buf = {}  # 缓冲 20 条后批量写入
        sources_buf = {}  # 缓冲触发时间来源

        def extract_one(bugid: str):
            """单个 bugid 的提取链路：视频 → 评论/描述 → 图片优先级"""
            if _stop_event.is_set():
                return bugid, "", "stopped", "已停止"
            try:
                issue = _jc.fetch_issue(bugid)
                if _stop_event.is_set():
                    return bugid, "", "stopped", "已停止"
                fields = issue.get("fields") or {}
                attachments = fields.get("attachment") or []
                try:
                    _gmt = _gm_times_of(attachments) or []
                    if _gmt:
                        _g0 = _gmt[0][0] if isinstance(_gmt[0], tuple) else _gmt[0]
                        _remember_ref_time(bugid, _g0.strftime("%Y-%m-%d %H:%M:%S") if hasattr(_g0, "strftime") else str(_g0)[:19])
                except Exception:
                    pass
                has_video = any((a.get("filename", "") or "").lower().endswith(_VIDEO_SUFFIXES) for a in attachments)
                tt, source = "", ""
                # 1. 视频 OCR（最高优先级）
                if has_video:
                    tt = _diag_extract_time_from_video_simple(issue, stop_event=_stop_event)
                    if tt:
                        source = "视频+gmlogger"
                    if _stop_event.is_set():
                        return bugid, "", "stopped", "已停止"
                # 2. 评论/描述/标题文本提取
                if not tt:
                    tt = _jc.extract_trigger_time_from_issue(issue, skip_attachment=True)
                    if tt:
                        source = "标准提取"
                    if _stop_event.is_set():
                        return bugid, "", "stopped", "已停止"
                # 3. 图片 OCR（最终兜底）
                if not tt and _img_headers:
                    img_tt, _ = _diag_extract_time_from_images(issue, _img_headers, img_to, bugid, stop_event=_stop_event)
                    if img_tt:
                        tt = img_tt
                        source = "图片OCR"
                if tt:
                    if not isinstance(tt, str):
                        tt = tt.strftime("%Y-%m-%d %H:%M:%S") if hasattr(tt, "strftime") else str(tt)
                    return bugid, tt, source, None
                return bugid, "", "failed", "优先级提取未获取到时间"
            except Exception as e:
                return bugid, "", "error", str(e)

        # 并行提取（max_workers=2）
        pool = concurrent.futures.ThreadPoolExecutor(max_workers=2)
        try:
            futures = {pool.submit(extract_one, b): b for b in to_extract}
            for future in concurrent.futures.as_completed(futures):
                if _stop_event.is_set():
                    # 取消未开始的任务
                    cancelled = sum(1 for f in futures if f.cancel())
                    # 把当前已完成的先批量写入，然后退出
                    if batch_buf:
                        _batch_save_trigger_times_to_cloud(batch_buf, sources_buf)
                        yield f"data: {_json.dumps({'type': 'batch_saved', 'count': len(batch_buf)}, ensure_ascii=False)}\n\n"
                        batch_buf.clear()
                        sources_buf.clear()
                    yield f"data: {_json.dumps({'type': 'stopped', 'current': done_count, 'total': total, 'hit': hit_count, 'message': f'已停止: {done_count}/{total} 完成, {hit_count} 命中, 取消{cancelled}个'}, ensure_ascii=False)}\n\n"
                    return
                try:
                    bugid, tt, source, error = future.result()
                except Exception as e:
                    bugid = futures[future]
                    tt, source, error = "", "error", str(e)

                # 跳过已停止的任务结果
                if source == "stopped":
                    continue

                done_count += 1
                if tt:
                    hit_count += 1
                    batch_buf[bugid] = tt
                    if source:
                        sources_buf[bugid] = source
                    # 缓冲 20 条后批量写入
                    if len(batch_buf) >= 20:
                        _batch_save_trigger_times_to_cloud(batch_buf, sources_buf)
                        yield f"data: {_json.dumps({'type': 'batch_saved', 'count': len(batch_buf)}, ensure_ascii=False)}\n\n"
                        batch_buf.clear()
                        sources_buf.clear()
                evt = {
                    'type': 'progress', 'current': done_count, 'total': total,
                    'hit': hit_count, 'bugid': bugid, 'time': tt or '',
                    'source': source, 'error': error or '',
                }
                yield f"data: {_json.dumps(evt, ensure_ascii=False)}\n\n"
        finally:
            pool.shutdown(wait=False, cancel_futures=True)

        # 余量写入
        if batch_buf:
            _batch_save_trigger_times_to_cloud(batch_buf, sources_buf)
            yield f"data: {_json.dumps({'type': 'batch_saved', 'count': len(batch_buf)}, ensure_ascii=False)}\n\n"
        done_msg = f'提取完成: {hit_count}/{total} 命中'
        if ai_skip_count:
            done_msg += f'，{ai_skip_count} 个已有AI分析跳过'
        yield f"data: {_json.dumps({'type': 'done', 'total': total, 'hit': hit_count, 'ai_skipped': ai_skip_count, 'message': done_msg}, ensure_ascii=False)}\n\n"

    # 注册作业：/api/jobs 可查进度，/api/jobs/{id}/stop 可单独停止本作业
    from src.core.job_registry import JOBS as _JOBS
    from src.api.routes.jobs import job_stream as _job_stream
    from src.api.routes.shared import batch_extract_stop_event as _job_stop
    _job = _JOBS.create("time_extract", "批量提取触发时间作业")
    return StreamingResponse(_job_stream(_job, event_generator, _job_stop), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})


@router.post("/api/trigger/batch_extract_time/stop", response_model=ApiResponse)
def stop_batch_extract_time():
    """停止批量提取触发时间"""
    from src.api.routes.shared import batch_extract_stop_event as _stop_event, ok
    _stop_event.set()
    logger.info("收到批量提取触发时间停止信号")
    return ok(None, "停止信号已发送，当前提取完成后停止")


@router.post("/api/trigger/compare_extract_time")
async def compare_extract_time(file: UploadFile = File(...)):
    """表格时间对比提取：上传 Excel/CSV，提取触发时间并与表格中的时间对比（SSE 流式）

    不写入云端，仅保存本地 CSV。
    """
    import json as _json
    import csv
    import re
    import io
    import concurrent.futures
    from datetime import datetime, timedelta
    from src.clients import jira_client as _jc
    from src.services.server_legacy import (
        _diag_extract_time_from_video_simple, _diag_extract_time_from_images,
        _VIDEO_SUFFIXES,
    )
    from src.config import load_config

    # 读取上传文件（带大小限制，避免超大表格打爆内存）
    from src.api.routes.shared import read_upload_limited as _read_upload_limited
    try:
        content = await _read_upload_limited(file, "上传表格")
    except ValueError as e:
        from src.api.routes.shared import fail as _fail
        return _fail(str(e))
    filename = file.filename or "upload.xlsx"
    _begin_attach_batch()  # 新批次：重置附件下载预算

    def event_generator():
        # ── Step 1: 解析表格 ──
        yield f"data: {_json.dumps({'type': 'parsing'}, ensure_ascii=False)}\n\n"
        rows_data = []
        col_names = []
        try:
            if filename.endswith('.csv'):
                text = content.decode('utf-8-sig', errors='replace')
                reader = csv.reader(io.StringIO(text))
                col_names = next(reader, [])
                for row in reader:
                    if any(cell.strip() for cell in row):
                        rows_data.append(row)
            else:
                import openpyxl
                wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
                ws = wb.active
                header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
                col_names = [str(c or '') for c in (header_row or [])]
                for row in ws.iter_rows(min_row=2, values_only=True):
                    vals = [str(c or '') if c is not None else '' for c in row]
                    if any(v.strip() for v in vals):
                        rows_data.append(vals)
                wb.close()
        except Exception as e:
            yield f"data: {_json.dumps({'type': 'error', 'message': f'文件解析失败: {e}'}, ensure_ascii=False)}\n\n"
            return

        # ── Step 2: 识别 Jira 号列和时间列 ──
        jira_col_idx = -1
        time_col_idx = -1
        jira_re = re.compile(r'[A-Z]+-\d+')
        time_keywords = ['问题发生时间', '触发时间', '时间锚点', '发生时间']
        jira_keywords = ['jira', 'Jira', 'JIRA', 'jira链接', 'Jira链接']

        # 按列名匹配
        for i, name in enumerate(col_names):
            if jira_col_idx < 0 and any(kw.lower() in name.lower() for kw in jira_keywords):
                jira_col_idx = i
            if time_col_idx < 0 and any(kw in name for kw in time_keywords):
                time_col_idx = i

        # 按内容匹配 Jira 列（兑底）
        if jira_col_idx < 0 and rows_data:
            for i in range(min(len(col_names), len(rows_data[0]))):
                if jira_re.search(str(rows_data[0][i] if i < len(rows_data[0]) else '')):
                    jira_col_idx = i
                    break

        if jira_col_idx < 0:
            yield f"data: {_json.dumps({'type': 'error', 'message': '未识别到 Jira 号列，请确保表格包含 Jira 链接列'}, ensure_ascii=False)}\n\n"
            return

        # 提取所有 bugid + 表格时间
        items = []
        for row in rows_data:
            if jira_col_idx >= len(row):
                continue
            cell = str(row[jira_col_idx])
            m = jira_re.search(cell)
            if not m:
                continue
            bugid = m.group(0)
            table_time = ""
            if time_col_idx >= 0 and time_col_idx < len(row):
                table_time = str(row[time_col_idx]).strip()
            items.append({"bugid": bugid, "table_time": table_time})

        if not items:
            yield f"data: {_json.dumps({'type': 'error', 'message': '未从表格中解析到有效 Jira 号'}, ensure_ascii=False)}\n\n"
            return

        cfg = load_config().get("jira_api", {})
        img_tok = cfg.get("token", "")
        img_to = max(int(cfg.get("timeout", 30)), 120)
        _img_headers = {}
        if img_tok:
            _img_headers["Authorization"] = f"Bearer {img_tok}"

        yield f"data: {_json.dumps({'type': 'start', 'total': len(items), 'jira_col': col_names[jira_col_idx] if jira_col_idx < len(col_names) else '', 'time_col': col_names[time_col_idx] if 0 <= time_col_idx < len(col_names) else ''}, ensure_ascii=False)}\n\n"

        # ── Step 3: 并行提取时间并对比 ──
        def extract_one(bugid: str):
            """提取单个 bugid 的触发时间"""
            try:
                issue = _jc.fetch_issue(bugid)
                fields = issue.get("fields") or {}
                attachments = fields.get("attachment") or []
                try:
                    _gmt = _gm_times_of(attachments) or []
                    if _gmt:
                        _g0 = _gmt[0][0] if isinstance(_gmt[0], tuple) else _gmt[0]
                        _remember_ref_time(bugid, _g0.strftime("%Y-%m-%d %H:%M:%S") if hasattr(_g0, "strftime") else str(_g0)[:19])
                except Exception:
                    pass
                has_video = any((a.get("filename", "") or "").lower().endswith(_VIDEO_SUFFIXES) for a in attachments)
                tt, source = "", ""
                if has_video:
                    tt = _diag_extract_time_from_video_simple(issue)
                    if tt:
                        source = "视频+gmlogger"
                if not tt:
                    tt = _jc.extract_trigger_time_from_issue(issue, skip_attachment=True)
                    if tt:
                        source = "标准提取"
                if not tt and _img_headers:
                    img_tt, _ = _diag_extract_time_from_images(issue, _img_headers, img_to, bugid)
                    if img_tt:
                        tt = img_tt
                        source = "图片OCR"
                if tt:
                    if not isinstance(tt, str):
                        tt = tt.strftime("%Y-%m-%d %H:%M:%S") if hasattr(tt, "strftime") else str(tt)
                    return bugid, tt, source, None
                return bugid, "", "failed", "未提取到时间"
            except Exception as e:
                return bugid, "", "error", str(e)

        def _parse_time(s: str):
            """尝试多种格式解析时间字符串"""
            if not s:
                return None
            s = s.strip()
            for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y/%m/%d %H:%M:%S",
                        "%Y/%m/%d %H:%M", "%Y.%m.%d %H:%M:%S", "%Y.%m.%d %H:%M"):
                try:
                    return datetime.strptime(s, fmt)
                except ValueError:
                    continue
            # 尝试 pandas Timestamp 或纯数字时间戳
            if s.isdigit() and len(s) >= 10:
                try:
                    ts = int(s)
                    if ts > 1e12:
                        ts = ts / 1000
                    return datetime.fromtimestamp(ts)
                except (ValueError, OSError):
                    pass
            return None

        done_count = 0
        matched = 0
        unmatched = 0
        failed = 0
        results_rows = []

        pool = concurrent.futures.ThreadPoolExecutor(max_workers=2)
        try:
            futures = {pool.submit(extract_one, it["bugid"]): it for it in items}
            for future in concurrent.futures.as_completed(futures):
                item = futures[future]
                bugid = item["bugid"]
                table_time_str = item["table_time"]
                try:
                    _, ext_time, source, error = future.result()
                except Exception as e:
                    ext_time, source, error = "", "error", str(e)

                done_count += 1
                # 对比逻辑
                match_result = ""
                remark = ""
                if not ext_time:
                    match_result = "失败"
                    remark = error or "未提取到时间"
                    failed += 1
                else:
                    ext_dt = _parse_time(ext_time)
                    tbl_dt = _parse_time(table_time_str)
                    if ext_dt and tbl_dt:
                        diff = abs((ext_dt - tbl_dt).total_seconds())
                        if diff <= 300:  # 5 分钟内视为一致
                            match_result = "成功"
                            matched += 1
                        else:
                            match_result = "不一致"
                            remark = f"差值 {int(diff)}s"
                            unmatched += 1
                    elif not tbl_dt:
                        match_result = "无法对比"
                        remark = "表格时间解析失败"
                        unmatched += 1
                    else:
                        match_result = "无法对比"
                        remark = "提取时间解析失败"
                        unmatched += 1

                results_rows.append({
                    "bugid": bugid, "table_time": table_time_str,
                    "extracted_time": ext_time, "match": match_result,
                    "source": source, "remark": remark,
                })
                evt = {
                    'type': 'progress', 'current': done_count, 'total': len(items),
                    'bugid': bugid, 'table_time': table_time_str,
                    'extracted_time': ext_time, 'match': match_result,
                    'source': source, 'remark': remark,
                }
                yield f"data: {_json.dumps(evt, ensure_ascii=False)}\n\n"
        finally:
            pool.shutdown(wait=False, cancel_futures=True)

        # ── Step 4: 保存本地 CSV ──
        ts_str = datetime.now().strftime("%Y%m%d_%H%M%S")
        csv_name = f"time_compare_{ts_str}.csv"
        csv_path = os.path.join(PROJECT_ROOT, "data", csv_name)
        os.makedirs(os.path.dirname(csv_path), exist_ok=True)
        try:
            with open(csv_path, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.writer(f)
                writer.writerow(["Jira号", "表格时间", "提取时间", "对比结果", "提取来源", "备注"])
                for r in results_rows:
                    writer.writerow([r["bugid"], r["table_time"], r["extracted_time"],
                                     r["match"], r["source"], r["remark"]])
        except Exception as e:
            logger.error("保存对比 CSV 失败: %s", e)
            csv_name = ""

        done_msg = f"对比完成: {matched} 成功, {unmatched} 不一致, {failed} 提取失败"
        yield f"data: {_json.dumps({'type': 'done', 'total': len(items), 'matched': matched, 'unmatched': unmatched, 'failed': failed, 'saved_path': csv_name, 'message': done_msg}, ensure_ascii=False)}\n\n"

    return StreamingResponse(event_generator(), media_type="text/event-stream",
                             headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"})
