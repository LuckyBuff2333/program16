"""从飞书导出的 Excel（数据沉淀.xlsx）导入 case_info 表

用法：
    python scripts/import_bitable_xlsx.py [xlsx路径]
    默认路径：~/Downloads/数据沉淀.xlsx
"""
import json
import os
import sys
from datetime import datetime

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, PROJECT_ROOT)

from src import db
from src.config import setup_logger

logger = setup_logger("import_xlsx")

# ── 字段映射：Excel 列名 → DB 字段名 ──
# 核心字段（直接映射为独立列）
CORE_FIELDS = {
    "jira号":           "case_id",
    "分析结果":         "analysis_result",
    "触发来源":         "trigger_source",
    "版本号":           "version",
    "分析问题时间":     "problem_time",
    "分析完成时间":     "finish_time",
    "分析重试次数":     "retry_count",
    "总耗时":           "total_seconds",
    "分析耗时":         "analysis_seconds",
    "错误类型":         "error_type",
    "错误信息":         "error_message",
    "AI分析问题所在模块": "ai_module",
    "问题所在实际模块": "actual_module",
    "分析日期":         "analysis_date",
}

# rootcause → 同时写入 rootcause 和 root_cause_summary
ROOT_CAUSE_FIELD = "rootcause"

# 需要转 float 的列
FLOAT_COLS = {"总耗时", "分析耗时", "access耗时", "reasoning耗时",
              "adjudication耗时", "reporting耗时", "结果置信度", "ADP置信度"}

# 需要转 int 的列
INT_COLS = {"分析重试次数"}

# 放入 extra_data 的列（非核心但保留）
EXTRA_FIELDS = {
    "版本变动": "version_change",
    "Jira链接": "jira_link",
    "人工审核结果": "audit_result",
    "access耗时": "access_duration",
    "reasoning耗时": "reasoning_duration",
    "adjudication耗时": "adjudication_duration",
    "reporting耗时": "reporting_duration",
    "AI评论总结": "ai_comment_summary",
    "置信度区间": "conf_interval",
    "分析时长区间": "duration_interval",
    "错误分类": "error_category",
    "bug状态": "bug_status",
    "ADP平台报告": "adp_report",
    "问题发生时间": "problem_occur_time",
    "ADP平台运行时间": "adp_runtime",
    "ADP现象级归因": "adp_attribution",
    "ADP置信度": "adp_confidence",
    "问题模块": "problem_module",
    "AI分析问题根因": "ai_root_cause_detail",
    "AI分析结果(飞书链接)": "ai_result_link",
}


def _safe_float(val):
    """安全转换浮点数"""
    if val is None or val == "" or val == "-":
        return None
    try:
        return float(val)
    except (ValueError, TypeError):
        return None


def _safe_int(val):
    """安全转换整数"""
    if val is None or val == "" or val == "-":
        return None
    try:
        return int(float(val))
    except (ValueError, TypeError):
        return None


def _safe_str(val):
    """安全转换字符串，空值返回 None"""
    if val is None:
        return None
    s = str(val).strip()
    return s if s and s not in ("-", "None", "nan", "") else None


def _seconds_to_duration(seconds):
    """秒数转可读时长字符串"""
    if seconds is None:
        return None
    try:
        s = float(seconds)
        if s < 60:
            return f"{s:.1f}s"
        elif s < 3600:
            return f"{s/60:.1f}min"
        else:
            h, rem = divmod(s, 3600)
            m = rem / 60
            return f"{h:.0f}h{m:.0f}min"
    except (ValueError, TypeError):
        return str(seconds)


def _map_analysis_result(result_str):
    """分析结果 → case_status"""
    if not result_str:
        return "历史导入"
    if result_str == "成功":
        return "已分析"
    elif result_str == "失败":
        return "分析失败"
    return result_str


def import_xlsx(xlsx_path: str, update_existing: bool = True) -> dict:
    """从 xlsx 导入数据到 case_info 表

    :param xlsx_path: xlsx 文件路径
    :param update_existing: True=更新已有记录, False=仅新增
    :return: 统计信息
    """
    import openpyxl

    if not os.path.isfile(xlsx_path):
        return {"error": f"文件不存在: {xlsx_path}"}

    print(f"[1/4] 读取 Excel: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active  # 取第一个 sheet

    # 读取表头
    headers = []
    for cell in next(ws.iter_rows(min_row=1, max_row=1)):
        headers.append(str(cell.value or "").strip())
    print(f"  表头 ({len(headers)} 列): {headers[:8]}...")

    # 读取数据行
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        row_dict = {}
        for i, val in enumerate(row):
            if i < len(headers) and headers[i]:
                row_dict[headers[i]] = val
        if row_dict.get("jira号"):
            rows.append(row_dict)
    wb.close()
    print(f"  数据行: {len(rows)} 条")

    # ── 2. 字段映射 ──
    print(f"\n[2/4] 字段映射...")
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    records = []
    for row in rows:
        jira = _safe_str(row.get("jira号"))
        if not jira:
            continue

        rec = {
            "case_id": jira,
            "vcu_number": jira,
            "created_at": now,
            "updated_at": now,
            "is_deleted": False,
            "converged": False,
            "decision_chain": "[]",
            "hypothesis_table": "[]",
        }

        # 核心字段映射
        for excel_col, db_field in CORE_FIELDS.items():
            val = row.get(excel_col)
            if excel_col in FLOAT_COLS:
                rec[db_field] = _safe_float(val)
            elif excel_col in INT_COLS:
                rec[db_field] = _safe_int(val)
            else:
                rec[db_field] = _safe_str(val)

        # 分析结果 → case_status
        rec["case_status"] = _map_analysis_result(_safe_str(row.get("分析结果")))

        # 置信度
        conf_val = _safe_float(row.get("结果置信度"))
        if conf_val is not None:
            # 如果值 > 1，说明是百分制，转为 0~1
            rec["conf"] = conf_val / 100 if conf_val > 1 else conf_val

        # 总耗时 → duration (可读格式)
        total_sec = rec.get("total_seconds")
        rec["duration"] = _seconds_to_duration(total_sec) if total_sec else _safe_str(row.get("总耗时"))

        # rootcause → root_cause_summary + rootcause
        rootcause_val = _safe_str(row.get("rootcause"))
        if rootcause_val:
            rec["rootcause"] = rootcause_val
            # root_cause_summary 优先用 rootcause 字段，如果为空则用 AI分析问题根因
            rec["root_cause_summary"] = rootcause_val
        else:
            ai_root = _safe_str(row.get("AI分析问题根因"))
            if ai_root:
                rec["root_cause_summary"] = ai_root[:500]  # 截断长文本

        # 问题描述：合并错误类型 + 错误信息
        err_type = _safe_str(row.get("错误类型"))
        err_msg = _safe_str(row.get("错误信息"))
        if err_type or err_msg:
            parts = []
            if err_type:
                parts.append(f"[{err_type}]")
            if err_msg:
                parts.append(err_msg)
            rec["problem_desc"] = " ".join(parts)

        # extra_data: 其他扩展字段
        extra = {}
        for excel_col, extra_key in EXTRA_FIELDS.items():
            val = row.get(excel_col)
            str_val = _safe_str(val)
            if str_val:
                if excel_col in FLOAT_COLS:
                    extra[extra_key] = _safe_float(val)
                else:
                    extra[extra_key] = str_val
        if extra:
            rec["extra_data"] = json.dumps(extra, ensure_ascii=False)

        records.append(rec)

    print(f"  有效记录: {len(records)} 条")

    # ── 3. 写入数据库 ──
    print(f"\n[3/4] 写入数据库...")
    session = db._get_session()
    try:
        existing = {r.case_id: r for r in session.query(db.CaseInfo).filter(
            db.CaseInfo.is_deleted == False).all()}
        existing_ids = set(existing.keys())

        new_records = []
        updated_count = 0

        for rec in records:
            cid = rec["case_id"]
            if cid not in existing_ids:
                new_records.append(rec)
            elif update_existing:
                # 更新已有记录
                obj = existing[cid]
                for key, val in rec.items():
                    if key in ("id", "case_id", "created_at"):
                        continue  # 不修改主键和创建时间
                    if val is not None:
                        setattr(obj, key, val)
                obj.updated_at = now
                updated_count += 1

        for rec in new_records:
            session.add(db.CaseInfo(**rec))
        session.commit()

        print(f"  新增: {len(new_records)} 条")
        print(f"  更新: {updated_count} 条")
        print(f"  跳过: {len(records) - len(new_records) - updated_count} 条")

        # ── 4. 统计 ──
        print(f"\n[4/4] 导入统计")
        total_in_db = session.query(db.CaseInfo).filter(
            db.CaseInfo.is_deleted == False).count()
        success_count = session.query(db.CaseInfo).filter(
            db.CaseInfo.is_deleted == False,
            db.CaseInfo.analysis_result == "成功").count()
        fail_count = session.query(db.CaseInfo).filter(
            db.CaseInfo.is_deleted == False,
            db.CaseInfo.analysis_result == "失败").count()
        print(f"  数据库总计: {total_in_db} 条")
        print(f"  分析成功: {success_count} 条")
        print(f"  分析失败: {fail_count} 条")

        return {
            "total": len(records),
            "inserted": len(new_records),
            "updated": updated_count,
            "skipped": len(records) - len(new_records) - updated_count,
            "db_total": total_in_db,
        }
    except Exception as e:
        session.rollback()
        logger.error("xlsx 导入失败: %s", e)
        print(f"  错误: {e}")
        return {"error": str(e)}
    finally:
        session.close()


if __name__ == "__main__":
    # 默认路径
    default_path = os.path.expanduser("~/Downloads/数据沉淀.xlsx")
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else default_path

    if not os.path.isfile(xlsx_path):
        print(f"文件不存在: {xlsx_path}")
        print(f"用法: python scripts/import_bitable_xlsx.py [xlsx路径]")
        sys.exit(1)

    print("=" * 60)
    print("  飞书数据沉淀表 → 本地数据库导入")
    print("=" * 60)
    result = import_xlsx(xlsx_path)
    print(f"\n完成: {json.dumps(result, ensure_ascii=False)}")
